"""Geração do Excel exibido na comparação de fundos."""

from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


COR_CABECALHO = "121212"
COR_TEXTO_CABECALHO = "FFFFFF"
COR_LINHA_ALTERNADA = "F3F4F6"
COR_POSITIVO = "008F5A"
COR_NEGATIVO = "D92D20"


def _rotulo_periodo(periodo: Dict[str, Any], indice: int) -> str:
    """Monta um rótulo legível para períodos padrão ou personalizados."""
    label = str(periodo.get("label") or "").strip()
    if label:
        return label.replace("<br>", " ").replace("\n", " ")

    data_inicial = periodo.get("data_inicial")
    data_final = periodo.get("data_final")
    if data_inicial and data_final:
        return f"{data_inicial} a {data_final}"
    return f"Período {indice}"


def gerar_excel_comparacao(
    fundos: List[Dict[str, Any]],
    periodos: List[Dict[str, Any]],
) -> BytesIO:
    """Exporta exatamente os resultados calculados e exibidos no comparativo."""
    if not fundos:
        raise ValueError("Nenhum fundo informado para exportação")
    if not periodos:
        raise ValueError("Nenhum período informado para exportação")

    wb = Workbook()
    ws = wb.active
    ws.title = "Comparativo de Fundos"

    ws["A1"] = "Comparativo de Fundos"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="00A86B")
    ws["A2"] = f"Gerado em: {datetime.now():%d/%m/%Y %H:%M}"
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")

    linha_cabecalho = 4
    cabecalho = ["CNPJ", "Nome do Fundo", "Subclasse"]
    for indice, periodo in enumerate(periodos, start=1):
        rotulo = _rotulo_periodo(periodo, indice)
        cabecalho.extend([
            f"Rentabilidade {rotulo} (%)",
            f"Volatilidade {rotulo} (%)",
        ])

    for coluna, titulo in enumerate(cabecalho, start=1):
        celula = ws.cell(row=linha_cabecalho, column=coluna, value=titulo)
        celula.font = Font(name="Arial", size=10, bold=True, color=COR_TEXTO_CABECALHO)
        celula.fill = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for offset, fundo in enumerate(fundos):
        linha = linha_cabecalho + 1 + offset
        valores = [
            fundo.get("cnpj", ""),
            fundo.get("nome", ""),
            fundo.get("id_subclasse") or "",
        ]
        variacoes = fundo.get("variacoes") or []

        for indice in range(len(periodos)):
            variacao = variacoes[indice] if indice < len(variacoes) else {}
            valores.extend([
                variacao.get("variacao_percentual", ""),
                variacao.get("volatilidade_percentual", ""),
            ])

        for coluna, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=coluna, value=valor)
            celula.font = Font(name="Arial", size=10)
            celula.alignment = Alignment(
                horizontal="left" if coluna <= 3 else "center",
                vertical="center",
            )
            if offset % 2 == 1:
                celula.fill = PatternFill(
                    start_color=COR_LINHA_ALTERNADA,
                    end_color=COR_LINHA_ALTERNADA,
                    fill_type="solid",
                )

        for coluna in range(4, len(cabecalho) + 1):
            celula = ws.cell(row=linha, column=coluna)
            if isinstance(celula.value, (int, float)):
                celula.number_format = "0.00"
                if coluna % 2 == 0:
                    celula.font = Font(
                        name="Arial",
                        size=10,
                        bold=True,
                        color=COR_POSITIVO if celula.value >= 0 else COR_NEGATIVO,
                    )

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 20
    for coluna in range(4, len(cabecalho) + 1):
        ws.column_dimensions[get_column_letter(coluna)].width = 22

    ultima_linha = linha_cabecalho + len(fundos)
    ultima_coluna = get_column_letter(len(cabecalho))
    ws.auto_filter.ref = f"A{linha_cabecalho}:{ultima_coluna}{ultima_linha}"
    ws.freeze_panes = f"D{linha_cabecalho + 1}"
    ws.row_dimensions[linha_cabecalho].height = 42

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
