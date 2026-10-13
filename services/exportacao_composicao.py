"""Exportadores da composição de cotas de fundos para Excel e PDF."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _validar(resultados: list[dict]) -> None:
    if not resultados:
        raise ValueError("Nenhuma composição informada para exportação")


def gerar_excel_composicao(resultados: list[dict]) -> BytesIO:
    _validar(resultados)
    wb = Workbook()
    ws = wb.active
    ws.title = "Composição de Fundos"
    ws.append(["Composição de Fundos"])
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="00A86B")
    ws.append([f"Gerado em: {datetime.now():%d/%m/%Y %H:%M}"])
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")
    ws.append([])
    cabecalho = ["Fundo consultado", "CNPJ do fundo", "Competência visível", "Fundo investido / ajuste", "CNPJ investido", "Subclasse", "Quantidade", "Valor de mercado (R$)", "% do PL do FOF"]
    for coluna, titulo in enumerate(cabecalho, 1):
        celula = ws.cell(row=4, column=coluna, value=titulo)
        celula.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="121212")
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    linha = 5
    for fundo in resultados:
        competencia = fundo.get("competencia") or "Sem composição aberta"
        linhas_resumo = [
            ("PL total do FOF", fundo.get("pl_total")),
            ("Soma das posições em fundos", fundo.get("soma_posicoes")),
            ("PL/POSIÇÕES - 1", fundo.get("percentual_pl_posicoes")),
        ]
        if fundo.get("pl_total") is not None:
            for rotulo, valor in linhas_resumo:
                valores = [fundo.get("nome", ""), fundo.get("cnpj", ""), competencia, rotulo, "", "", None, valor, None]
                if rotulo == "PL/POSIÇÕES - 1":
                    valores[7], valores[8] = None, valor
                for coluna, conteudo in enumerate(valores, 1):
                    celula = ws.cell(row=linha, column=coluna, value=conteudo)
                    celula.font = Font(name="Arial", size=10, bold=True)
                    celula.fill = PatternFill("solid", fgColor="E8EAED")
                    if coluna == 8 and isinstance(conteudo, (int, float)):
                        celula.number_format = "#,##0.00"
                    elif coluna == 9 and isinstance(conteudo, (int, float)):
                        celula.number_format = '0.00"%"'
                linha += 1

        posicoes = fundo.get("posicoes") or [{}]
        for posicao in posicoes:
            valores = [
                fundo.get("nome", ""), fundo.get("cnpj", ""), competencia,
                posicao.get("nome", ""), posicao.get("cnpj", ""), posicao.get("id_subclasse", ""),
                posicao.get("quantidade"), posicao.get("valor"), posicao.get("percentual_pl"),
            ]
            for coluna, valor in enumerate(valores, 1):
                celula = ws.cell(row=linha, column=coluna, value=valor)
                celula.font = Font(name="Arial", size=10)
                if linha % 2 == 0:
                    celula.fill = PatternFill("solid", fgColor="F3F4F6")
                if coluna in (7, 8) and isinstance(valor, (int, float)):
                    celula.number_format = "#,##0.00"
                elif coluna == 9 and isinstance(valor, (int, float)):
                    celula.number_format = '0.00"%"'
            linha += 1
        if not fundo.get("posicoes") and fundo.get("sigilo"):
            sigilo = fundo["sigilo"]
            prazo = f"; prazo até {sigilo['ate']}" if sigilo.get("ate") else ""
            ws.cell(row=linha - 1, column=3, value=f"Sob sigilo em {sigilo.get('competencia') or 'competência recente'}{prazo}")

    larguras = [42, 22, 22, 64, 22, 24, 18, 24, 18]
    for coluna, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(coluna)].width = largura
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:I{max(4, linha - 1)}"
    ws.row_dimensions[4].height = 34
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def gerar_pdf_composicao(resultados: list[dict]) -> BytesIO:
    from services.pdf_composicao import montar_pdf

    _validar(resultados)
    return montar_pdf(resultados)
